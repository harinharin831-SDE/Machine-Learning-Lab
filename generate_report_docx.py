import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, color="D1D5DB", sz="4", val="single"):
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def add_code_box(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_border(cell, color="D1D5DB", sz="6")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Remove default paragraph in cell or write to it
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    lines = code_text.strip().split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(30, 30, 30)
    
    # Add a little space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_figure(doc, image_path, caption_text, width_inches=5.8):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(10)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(40, 40, 40)
    else:
        print(f"Warning: Image not found: {image_path}")

def build_project_document(output_docx_path):
    doc = docx.Document()
    
    # Page Setup - 1 inch margins all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # -------------------------------------------------------------
    # 1. TITLE & SUBTITLE
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(3)
    run_title = p_title.add_run("Kidney Disease Detection and Classification from Ultrasound Images using Machine Learning and Hybrid Feature Extraction")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("(Kidney Ultrasound Image Classification using Machine Learning & Computer Vision)")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    
    # -------------------------------------------------------------
    # 2. PROBLEM STATEMENT
    # -------------------------------------------------------------
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(10)
    p_h1.paragraph_format.space_after = Pt(4)
    r = p_h1.add_run("Problem Statement:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    p_ps1 = doc.add_paragraph()
    p_ps1.paragraph_format.line_spacing = 1.15
    p_ps1.paragraph_format.space_after = Pt(6)
    p_ps1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ps1.add_run(
        "Kidney diseases, including renal cysts, nephrolithiasis (kidney stones), and renal neoplasms (tumors), represent major causes of chronic kidney damage and morbidity worldwide. Ultrasound is the premier frontline imaging modality for urinary tract assessment due to its non-invasive nature, lack of ionizing radiation, low cost, and real-time acquisition. However, ultrasound images suffer from low signal-to-noise ratios, acoustic speckle noise, acoustic shadows, and operator-dependent scanning artifacts. Manual screening of kidney ultrasound scans by radiologists is accurate but time-consuming, subjective, prone to fatigue and inter-observer variability, and not easily scalable to large populations, especially in resource-constrained rural clinics."
    )
    
    p_ps2 = doc.add_paragraph()
    p_ps2.paragraph_format.line_spacing = 1.15
    p_ps2.paragraph_format.space_after = Pt(12)
    p_ps2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ps2.add_run(
        "There is, therefore, a need for an automated, fast, accessible, and transparent system that can analyze a kidney ultrasound scan and classify it into diagnostic categories (Normal, Cyst, Stone, Tumor) with a measurable confidence score, so that suspected pathologies can be flagged early for a specialist's review. This project designs and implements a web-based Kidney Image Classification system that performs a 5-stage digital image preprocessing pipeline, hybrid handcrafted feature extraction (Histogram of Oriented Gradients and Gray-Level Co-occurrence Matrix), and classification using traditional Machine Learning models (Random Forest, Support Vector Machine, and K-Nearest Neighbors), and evaluates the trained models using standard classification metrics (Accuracy, Precision, Recall, F1-Score, Confusion Matrix)."
    )
    
    # -------------------------------------------------------------
    # 3. OBJECTIVES
    # -------------------------------------------------------------
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(8)
    p_h2.paragraph_format.space_after = Pt(4)
    r = p_h2.add_run("Objectives:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    objectives = [
        "To collect and preprocess a labelled dataset of Normal, Cyst, Stone, and Tumor kidney ultrasound images.",
        "To develop an automated 5-stage digital image preprocessing pipeline (spatial resizing, grayscale conversion, Gaussian blur denoising, histogram equalization, and min-max normalization).",
        "To extract an 8,160-dimensional hybrid handcrafted feature vector combining Histogram of Oriented Gradients (HOG) structural edge orientations and Gray-Level Co-occurrence Matrix (GLCM) second-order statistical texture properties from each image.",
        "To design, train, evaluate, and save traditional Machine Learning classifiers (Random Forest, Support Vector Machine with RBF kernel, and K-Nearest Neighbors) using a standardized Scikit-Learn pipeline.",
        "To build an interactive Streamlit web front-end that lets a user explore datasets, inspect intermediate preprocessing stages, visualize extracted features, and upload ultrasound images to view the real-time diagnosis.",
        "To evaluate the trained model on an unseen test set and report Accuracy, Precision, Recall, F1-score and the Confusion Matrix.",
        "To persist every prediction (filename, predicted class, confidence, timestamp) and provide downloadable evaluation reports for record-keeping."
    ]
    
    for i, obj in enumerate(objectives, 1):
        p_obj = doc.add_paragraph()
        p_obj.paragraph_format.line_spacing = 1.15
        p_obj.paragraph_format.space_after = Pt(3)
        p_obj.paragraph_format.left_indent = Inches(0.3)
        r_num = p_obj.add_run(f"{i}. ")
        r_num.font.bold = True
        p_obj.add_run(obj)
        
    # -------------------------------------------------------------
    # 4. SOFTWARE / HARDWARE REQUIRED
    # -------------------------------------------------------------
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(10)
    p_h3.paragraph_format.space_after = Pt(4)
    r = p_h3.add_run("Software / Hardware Required:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    requirements = [
        ("Operating System", "Windows / Linux / macOS"),
        ("IDE", "Visual Studio Code"),
        ("Language", "Python 3.x (Python 3.10+)"),
        ("Frontend / Web Dashboard", "Streamlit"),
        ("Machine Learning", "scikit-learn (RandomForestClassifier, SVC, KNeighborsClassifier, StandardScaler, Pipeline)"),
        ("Image Processing", "OpenCV (cv2), Pillow (PIL), NumPy"),
        ("Feature Extraction", "scikit-image (hog, graycomatrix, graycoprops)"),
        ("Data Analysis & Visualization", "Pandas, Matplotlib, Seaborn, SciPy"),
        ("Model Persistence", "Joblib"),
        ("Dataset", "Multi-Class Kidney Ultrasound Dataset (Normal, Cyst, Stone, Tumor classes)"),
        ("Version Control", "Git")
    ]
    
    for item, desc in requirements:
        p_req = doc.add_paragraph()
        p_req.paragraph_format.space_after = Pt(2.5)
        p_req.paragraph_format.left_indent = Inches(0.2)
        r_bullet = p_req.add_run("• ")
        r_bullet.font.bold = True
        r_item = p_req.add_run(f"{item}: ")
        r_item.font.bold = True
        p_req.add_run(desc)
        
    # -------------------------------------------------------------
    # 5. PROJECT DIRECTORY STRUCTURE
    # -------------------------------------------------------------
    p_h4 = doc.add_paragraph()
    p_h4.paragraph_format.space_before = Pt(12)
    p_h4.paragraph_format.space_after = Pt(4)
    r = p_h4.add_run("Project Directory Structure:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    p_dir_desc = doc.add_paragraph()
    p_dir_desc.paragraph_format.line_spacing = 1.15
    p_dir_desc.paragraph_format.space_after = Pt(6)
    p_dir_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dir_desc.add_run(
        "The project (\"KidneyDiseaseClassifier\") is organized as a modular Machine Learning pipeline with separate dataset, feature extraction, trained-model, preprocessing, and interactive Streamlit web dashboard components, as shown below."
    )
    
    add_figure(doc, "doc_figures/fig0_directory_structure.png", "Fig 1: Project directory structure of KidneyDiseaseClassifier in Visual Studio Code", width_inches=4.8)
    
    # -------------------------------------------------------------
    # 6. PROCEDURE
    # -------------------------------------------------------------
    p_h5 = doc.add_paragraph()
    p_h5.paragraph_format.space_before = Pt(12)
    p_h5.paragraph_format.space_after = Pt(4)
    r = p_h5.add_run("Procedure:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    procedures = [
        "Collect a labelled kidney ultrasound dataset (Normal, Cyst, Stone, Tumor) and organize it into dataset/ class directories with train and test partitions.",
        "Launch the Streamlit web application (app.py) or execute train_model.py to initialize the automated pipeline.",
        "Browse and inspect dataset distributions, class balances, and high-resolution ultrasound scans through the Dataset Explorer (Step 1).",
        "Preprocess the image — resize to 128×128, convert to single-channel luminance grayscale, apply Gaussian blur (5×5 kernel, σ=1.0) for speckle noise suppression, apply histogram equalization (CDF stretching) for tissue contrast enhancement, and normalize pixel values to [0.0, 1.0] (Step 2).",
        "Extract an 8,160-dimensional hybrid handcrafted feature vector combining 8,100 HOG (Histogram of Oriented Gradients) edge-orientation descriptors and 60 GLCM (Gray-Level Co-occurrence Matrix) second-order statistical texture properties (Step 3).",
        "Partition the dataset into 80% training and 20% test splits using stratified sampling to preserve relative class frequencies across splits.",
        "Train classical Machine Learning classifiers (Random Forest, Support Vector Machine with RBF kernel, and K-Nearest Neighbors) using a Scikit-Learn StandardScaler pipeline, and save the trained model bundle to models/kidney_model.pkl using Joblib (Step 4).",
        "Feed the extracted feature vector of an uploaded ultrasound image to the trained pipeline to obtain the predicted class (Normal / Cyst / Stone / Tumor) along with a multi-class confidence probability percentage (Step 5).",
        "Evaluate the trained model on the held-out test set and compute Accuracy, Weighted Precision, Recall, F1-Score, and the Confusion Matrix Heatmap (Step 6).",
        "Export publication-ready evaluation artifacts (outputs/confusion_matrix.png, outputs/classification_report.csv) and persist diagnostic predictions with timestamps (Step 7)."
    ]
    
    for i, proc in enumerate(procedures, 1):
        p_pr = doc.add_paragraph()
        p_pr.paragraph_format.line_spacing = 1.15
        p_pr.paragraph_format.space_after = Pt(3.5)
        p_pr.paragraph_format.left_indent = Inches(0.3)
        r_num = p_pr.add_run(f"{i}. ")
        r_num.font.bold = True
        p_pr.add_run(proc)
        
    # -------------------------------------------------------------
    # 7. IMPLEMENTATION CODE
    # -------------------------------------------------------------
    p_h6 = doc.add_paragraph()
    p_h6.paragraph_format.space_before = Pt(14)
    p_h6.paragraph_format.space_after = Pt(4)
    r = p_h6.add_run("Implementation Code:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    p_code_intro = doc.add_paragraph()
    p_code_intro.paragraph_format.line_spacing = 1.15
    p_code_intro.paragraph_format.space_after = Pt(6)
    p_code_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_ci = p_code_intro.add_run("Representative implementation of the key modules used in the project, based on the pipeline shown in the application (adapt variable/file names to match your exact source files).")
    r_ci.font.italic = True
    
    # 7.1 Image Preprocessing
    p_c1_title = doc.add_paragraph()
    p_c1_title.paragraph_format.space_before = Pt(8)
    p_c1_title.paragraph_format.space_after = Pt(3)
    r = p_c1_title.add_run("1. Image Preprocessing (preprocessing.py)")
    r.font.bold = True
    
    code_prep = """import cv2
import numpy as np

def preprocess_image(image_path, target_size=(128, 128)):
    # 1. Load Original RGB Image
    img = cv2.imread(image_path)
    img_rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)           # RGB conversion
    
    # 2. Resize to 128x128 standard spatial resolution
    resized = cv2.resize(img_rgb, target_size, interpolation=cv2.INTER_LINEAR)
    
    # 3. Grayscale conversion (Luminance Y = 0.299R + 0.587G + 0.114B)
    gray = cv2.cvtColor(resized, cv2.COLOR_RGB2GRAY)
    
    # 4. Noise Reduction using Gaussian Blur filter
    denoised = cv2.GaussianBlur(gray, (5, 5), 1.0)           # Speckle filtering
    
    # 5. Contrast Enhancement using Histogram Equalization
    enhanced = cv2.equalizeHist(denoised)                    # CDF stretch
    
    # 6. Min-Max Intensity Normalization to [0.0, 1.0]
    normalized = enhanced.astype(np.float32) / 255.0         # Normalization
    
    return {
        "resized": resized,
        "gray": gray,
        "denoised": denoised,
        "enhanced": enhanced,
        "normalized": normalized,
        "final": enhanced                                    # uint8 for HOG/GLCM
    }"""
    add_code_box(doc, code_prep)
    
    # 7.2 Feature Extraction
    p_c2_title = doc.add_paragraph()
    p_c2_title.paragraph_format.space_before = Pt(8)
    p_c2_title.paragraph_format.space_after = Pt(3)
    r = p_c2_title.add_run("2. Feature Extraction (feature_extraction.py)")
    r.font.bold = True
    
    code_feat = """import numpy as np
from skimage.feature import hog, graycomatrix, graycoprops

def extract_features(img_gray):
    # 1. Histogram of Oriented Gradients (HOG) - 8,100 features
    hog_feats = hog(
        img_gray,
        orientations=9,
        pixels_per_cell=(8, 8),
        cells_per_block=(2, 2),
        block_norm='L2-Hys',
        transform_sqrt=True
    )  # shape: (8100,)
    
    # 2. Gray-Level Co-occurrence Matrix (GLCM) - 60 features
    distances = [1, 2, 3]
    angles = [0, np.pi/4, np.pi/2, 3*np.pi/4]               # 0, 45, 90, 135 deg
    glcm = graycomatrix(img_gray, distances=distances, angles=angles,
                        levels=256, symmetric=True, normed=True)
    
    glcm_props = ['contrast', 'dissimilarity', 'homogeneity', 'energy', 'correlation']
    glcm_feats = []
    for prop in glcm_props:
        matrix = graycoprops(glcm, prop)
        glcm_feats.extend(matrix.flatten())
    glcm_feats = np.array(glcm_feats, dtype=np.float32)     # shape: (60,)
    
    # 3. Concatenate into combined 8,160-dimensional feature vector
    combined_vector = np.hstack((hog_feats, glcm_feats)).astype(np.float32)
    return combined_vector  # 8,160-dimensional feature vector"""
    add_code_box(doc, code_feat)
    
    # 7.3 Model Training
    p_c3_title = doc.add_paragraph()
    p_c3_title.paragraph_format.space_before = Pt(8)
    p_c3_title.paragraph_format.space_after = Pt(3)
    r = p_c3_title.add_run("3. Model Training (model_training.py / train_model.py)")
    r.font.bold = True
    
    code_train = """from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.ensemble import RandomForestClassifier
from sklearn.svm import SVC
from sklearn.model_selection import train_test_split
import joblib, numpy as np

# Load extracted feature matrix and class labels
X = np.load('processed/features.npy')                       # (N, 8160)
y = np.load('processed/labels.npy')                         # 0=Cyst, 1=Normal, 2=Stone, 3=Tumor
classes = ['Cyst', 'Normal', 'Stone', 'Tumor']

# Stratified Train-Test Split (80% Train, 20% Test)
X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.20, stratify=y, random_state=42
)

# Build Scikit-Learn Pipeline: StandardScaler + RandomForestClassifier
pipeline = Pipeline([
    ('scaler', StandardScaler()),
    ('classifier', RandomForestClassifier(n_estimators=100, max_depth=None, random_state=42))
])

# Train model on 8,160-D training feature vectors
pipeline.fit(X_train, y_train)

# Evaluate validation accuracy
test_acc = pipeline.score(X_test, y_test)
print(f'Test Set Classification Accuracy: {test_acc * 100:.2f}%')

# Save model bundle
joblib.dump({'pipeline': pipeline, 'class_names': classes}, 'models/kidney_model.pkl')"""
    add_code_box(doc, code_train)
    
    # 7.4 Prediction Route
    p_c4_title = doc.add_paragraph()
    p_c4_title.paragraph_format.space_before = Pt(8)
    p_c4_title.paragraph_format.space_after = Pt(3)
    r = p_c4_title.add_run("4. Single Image Prediction Pipeline (prediction.py / app.py)")
    r.font.bold = True
    
    code_pred = """import joblib, numpy as np
from preprocessing import preprocess_image
from feature_extraction import extract_features

# Load trained pipeline bundle
model_package = joblib.load('models/kidney_model.pkl')
pipeline = model_package['pipeline']
CLASSES = model_package['class_names']

def predict_kidney_image(image_input):
    # 1. Preprocess ultrasound image (128x128 5-stage pipeline)
    prep = preprocess_image(image_input)
    
    # 2. Extract 8,160-D feature vector (HOG + GLCM)
    features = extract_features(prep['enhanced']).reshape(1, -1)
    
    # 3. Model prediction & probability distribution
    pred_idx = int(pipeline.predict(features)[0])
    probabilities = pipeline.predict_proba(features)[0]
    
    result = {
        'prediction': CLASSES[pred_idx],
        'confidence': round(float(probabilities[pred_idx]) * 100, 1),
        'class_probabilities': {CLASSES[i]: round(float(probabilities[i])*100, 1) 
                                for i in range(len(CLASSES))},
        'feature_count': features.shape[1]
    }
    return result"""
    add_code_box(doc, code_pred)
    
    # 7.5 Evaluation
    p_c5_title = doc.add_paragraph()
    p_c5_title.paragraph_format.space_before = Pt(8)
    p_c5_title.paragraph_format.space_after = Pt(3)
    r = p_c5_title.add_run("5. Performance Evaluation & Metrics Visualization (evaluation.py)")
    r.font.bold = True
    
    code_eval = """from sklearn.metrics import (
    accuracy_score, precision_score, recall_score, f1_score, confusion_matrix
)
import matplotlib.pyplot as plt
import seaborn as sns

def evaluate_model_performance(pipeline, X_test, y_test, class_names):
    # Model predictions on unseen test split
    y_pred = pipeline.predict(X_test)
    
    # Compute multi-class metrics (weighted)
    metrics = {
        'accuracy': accuracy_score(y_test, y_pred),
        'precision': precision_score(y_test, y_pred, average='weighted', zero_division=0),
        'recall': recall_score(y_test, y_pred, average='weighted', zero_division=0),
        'f1_score': f1_score(y_test, y_pred, average='weighted', zero_division=0)
    }
    
    # Multi-Class Confusion Matrix Heatmap
    cm = confusion_matrix(y_test, y_pred)
    plt.figure(figsize=(6, 5))
    sns.heatmap(cm, annot=True, fmt='d', cmap='Blues',
                xticklabels=class_names, yticklabels=class_names)
    plt.title('Test Set Confusion Matrix')
    plt.ylabel('True Class')
    plt.xlabel('Predicted Class')
    plt.savefig('outputs/confusion_matrix.png', dpi=300, bbox_inches='tight')
    plt.close()
    
    return metrics"""
    add_code_box(doc, code_eval)
    
    # -------------------------------------------------------------
    # 8. OUTPUT (STEP-BY-STEP SCREENSHOTS)
    # -------------------------------------------------------------
    p_h7 = doc.add_paragraph()
    p_h7.paragraph_format.space_before = Pt(14)
    p_h7.paragraph_format.space_after = Pt(6)
    r = p_h7.add_run("Output (Step-by-Step Screenshots):")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    # Figure 2: System Architecture
    add_figure(
        doc,
        "doc_figures/fig1_architecture.png",
        "Fig 2: End-to-End System Architecture: Input Ultrasound Image → 5-Stage Digital Preprocessing → Hybrid Feature Extraction (HOG + GLCM) → StandardScaler Pipeline → ML Classifiers (SVM, Random Forest, KNN) → Performance Evaluation & Inference",
        width_inches=5.8
    )
    
    # Figure 3: Load / Explore Dataset
    add_figure(
        doc,
        "doc_figures/fig5_class_distribution.png",
        "Fig 3: Step 1 — Load Input Image & Dataset Explorer: Multi-class ultrasound dataset explorer showing sample scans across Normal, Cyst, Stone, and Tumor categories",
        width_inches=5.8
    )
    
    # Figure 4: 5-Stage Preprocessing
    add_figure(
        doc,
        "doc_figures/fig2_preprocessing_stages.png",
        "Fig 4: Step 2 — 5-Stage Digital Image Preprocessing Lab: Step-by-step visual transformations (Stage 01 Raw Input, Stage 02 Resized & Gray, Stage 03 Gaussian Denoised, Stage 04 Hist. Equalized, Stage 05 Normalized)",
        width_inches=5.8
    )
    
    # Figure 5: Histogram Transformation
    add_figure(
        doc,
        "doc_figures/fig3_histograms.png",
        "Fig 5: Step 2 — Pixel Intensity Histogram Transformation: Grayscale distribution before equalization vs. uniform dynamic spread after histogram equalization",
        width_inches=5.8
    )
    
    # Figure 6: Feature Extraction Studio
    add_figure(
        doc,
        "doc_figures/fig4_hog_glcm.png",
        "Fig 6: Step 3 — Feature Extraction Studio: 8,100-dimensional HOG gradient orientation vector field and 60-dimensional GLCM statistical texture descriptors breakdown",
        width_inches=5.8
    )
    
    # Figure 7: Model Training Studio
    add_figure(
        doc,
        "doc_figures/fig6_model_comparison.png",
        "Fig 7: Step 4 — Model Training Studio: Multi-model benchmark performance comparison across Random Forest, SVM (RBF), and KNN classifiers",
        width_inches=5.8
    )
    
    # Figure 8: Image Classification Result
    add_figure(
        doc,
        "doc_figures/fig8_prediction_result.png",
        "Fig 8: Step 5 — Classification result: Diagnostic prediction (Normal Kidney Healthy / Renal Cyst) with confidence percentage and class probability distribution",
        width_inches=5.8
    )
    
    # Figure 9: Performance Evaluation Metrics & Confusion Matrix
    add_figure(
        doc,
        "doc_figures/fig7_confusion_matrix.png",
        "Fig 9: Step 6 — Performance Evaluation Metrics (Accuracy 91.67%, Weighted Precision 93.75%, Recall 91.67%, F1-Score 91.43%) and Confusion Matrix heatmap",
        width_inches=5.0
    )
    
    # -------------------------------------------------------------
    # 9. RESULT
    # -------------------------------------------------------------
    p_h8 = doc.add_paragraph()
    p_h8.paragraph_format.space_before = Pt(14)
    p_h8.paragraph_format.space_after = Pt(4)
    r = p_h8.add_run("Result:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    p_res = doc.add_paragraph()
    p_res.paragraph_format.line_spacing = 1.15
    p_res.paragraph_format.space_after = Pt(12)
    p_res.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_res.add_run(
        "A web-based Kidney Disease Image Classification system was successfully designed and implemented using a Streamlit frontend, a modular Python processing pipeline, classical Machine Learning classifiers (scikit-learn Random Forest, SVM, and KNN), and Joblib model serialization. Ultrasound images are preprocessed through a 5-stage digital pipeline, an 8,160-dimensional hybrid feature vector (8,100 HOG edge gradients + 60 GLCM statistical texture descriptors) is extracted, and the trained Random Forest model classifies the image as Normal, Cyst, Stone, or Tumor with a multi-class confidence score. On the held-out test set, the Random Forest model achieved an outstanding Accuracy of 91.67%, Weighted Precision of 93.75%, Weighted Recall of 91.67%, and an F1-Score of 91.43%, and every prediction is logged with diagnostic confidence probabilities and downloadable evaluation reports."
    )
    
    doc.save(output_docx_path)
    print(f"Document successfully created and saved to: {output_docx_path}")

if __name__ == "__main__":
    out_path = "Kidney_Disease_Classification_Project_Report.docx"
    build_project_document(out_path)
